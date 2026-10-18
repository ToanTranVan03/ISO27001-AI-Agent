from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from backend import main
from agent_core import graph_builder
from evidence_collectors.document_parser import extract_document_text


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Approved information security policy.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Security"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_parser_extracts_paragraphs_and_tables():
    text = extract_document_text("policy.docx", make_docx())

    assert "Approved information security policy." in text
    assert "Owner | Security" in text


def test_parser_rejects_unsupported_extension():
    try:
        extract_document_text("policy.txt", b"policy")
    except ValueError as exc:
        assert "Only PDF and DOCX" in str(exc)
    else:
        raise AssertionError("Unsupported file extension should be rejected.")


def test_upload_saves_extracted_document(monkeypatch):
    saved = {}
    monkeypatch.setattr(main.database, "initialize_database", lambda: None)
    monkeypatch.setattr(
        main.database,
        "get_control",
        lambda _: {"expected_evidence_types": ["EVID_DOC_POLICY"]},
    )
    monkeypatch.setattr(
        main.database,
        "save_evidence",
        lambda **values: saved.update(values),
    )

    with TestClient(main.app) as client:
        response = client.post(
            "/api/evidence",
            data={"control_id": "A.6.4", "evidence_type": "EVID_DOC_POLICY"},
            files={
                "file": (
                    "policy.docx",
                    make_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert response.status_code == 201
    assert saved["control_id"] == "A.6.4"
    assert "Approved information security policy." in saved["extracted_text"]
    assert len(saved["sha256_hash"]) == 64


def test_upload_rejects_unexpected_evidence_type(monkeypatch):
    monkeypatch.setattr(main.database, "initialize_database", lambda: None)
    monkeypatch.setattr(
        main.database,
        "get_control",
        lambda _: {"expected_evidence_types": ["EVID_DOC_POLICY"]},
    )

    with TestClient(main.app) as client:
        response = client.post(
            "/api/evidence",
            data={"control_id": "A.6.4", "evidence_type": "EVID_TECH_IAM_AUTH"},
            files={"file": ("policy.docx", make_docx())},
        )

    assert response.status_code == 422


def test_collect_node_reads_evidence_from_database(monkeypatch):
    monkeypatch.setattr(
        graph_builder,
        "collect_evidence",
        lambda _: (
            [{"evidence_id": "ev-1", "source": "policy.pdf", "content": "approved"}],
            {},
        ),
    )
    state = {
        "control_id": "A.6.4",
        "collected_evidences": [],
        "deterministic_rule": {},
        "audit_trail": [],
    }

    result = graph_builder.node_collect(state)

    assert result["collected_evidences"][0]["evidence_id"] == "ev-1"
    assert result["audit_trail"] == ["COLLECT"]


def test_home_page_serves_frontend(monkeypatch):
    monkeypatch.setattr(main.database, "initialize_database", lambda: None)

    with TestClient(main.app) as client:
        response = client.get("/")

    assert response.status_code == 200
    assert "ISO 27001 Evidence Evaluator" in response.text


def test_evaluation_uses_database_control_and_persists_result(monkeypatch):
    saved = {}
    monkeypatch.setattr(main.database, "initialize_database", lambda: None)
    monkeypatch.setattr(
        main.database,
        "get_control",
        lambda _: {
            "control_id": "A.6.4",
            "name": "Disciplinary process",
            "objective": "Published disciplinary process",
        },
    )
    monkeypatch.setattr(main.database, "save_evaluation", lambda result: saved.update(result))
    monkeypatch.setattr(
        main.agent_graph,
        "invoke",
        lambda state: {
            **state,
            "compliance_status": "Compliant",
            "confidence": 0.9,
            "rationale": "Policy evidence supports this control.",
            "collected_evidences": [
                {"evidence_id": "ev-1", "content": "Approved policy."}
            ],
            "audit_trail": ["PLAN", "COLLECT", "EVALUATE", "VERIFY", "REPORT"],
        },
    )

    with TestClient(main.app) as client:
        response = client.post("/api/evaluate", json={"control_id": "A.6.4"})

    assert response.status_code == 200
    assert response.json()["evidence_ids"] == ["ev-1"]
    assert saved["status"] == "Compliant"
