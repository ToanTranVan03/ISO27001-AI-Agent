"""
Document Evidence Parser
========================
Trích xuất văn bản từ file PDF và DOCX để nạp vào node COLLECT của AI Agent.

Thư viện ưu tiên:
- PDF  : pdfplumber (chính), pypdf (fallback)
- DOCX : python-docx
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

MAX_DOCUMENT_CHARS = 100_000


def _extract_pdf(content: bytes) -> str:
    """Trích xuất text từ PDF, thử pdfplumber trước rồi fallback sang pypdf."""
    # --- pdfplumber (ưu tiên, nhận diện bảng tốt hơn) ---
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(BytesIO(content)) as pdf:
            if not pdf.pages:
                raise ValueError("PDF không có trang nào.")
            pages: list[str] = []
            for page in pdf.pages:
                text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
                # Trích xuất thêm bảng nếu có
                for table in page.extract_tables() or []:
                    rows = []
                    for row in table:
                        rows.append(" | ".join(cell or "" for cell in row))
                    text += "\n" + "\n".join(rows)
                pages.append(text)
            return "\n".join(pages)

    except ImportError:
        pass  # pdfplumber chưa cài → thử pypdf
    except Exception:
        pass  # file lỗi → thử pypdf

    # --- pypdf (fallback) ---
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDF files are not supported.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    """Trích xuất text từ DOCX (đoạn văn + bảng)."""
    from docx import Document  # type: ignore

    document = Document(BytesIO(content))
    paragraphs = [p.text for p in document.paragraphs]
    tables = [
        "\n".join(
            " | ".join(cell.text for cell in row.cells)
            for row in table.rows
        )
        for table in document.tables
    ]
    return "\n".join(paragraphs + tables)


def extract_document_text(filename: str, content: bytes) -> str:
    """
    Public API — nhận filename và raw bytes, trả về văn bản thuần.

    Args:
        filename: Tên file (dùng để xác định phần mở rộng).
        content:  Nội dung file dạng bytes.

    Returns:
        Chuỗi văn bản đã trích xuất (không quá MAX_DOCUMENT_CHARS ký tự).

    Raises:
        ValueError: Nếu định dạng không được hỗ trợ, file rỗng, hoặc quá lớn.
    """
    extension = Path(filename).suffix.lower()

    try:
        if extension == ".pdf":
            text = _extract_pdf(content)
        elif extension == ".docx":
            text = _extract_docx(content)
        else:
            raise ValueError(
                f"Định dạng '{extension}' không được hỗ trợ. Chỉ chấp nhận PDF và DOCX."
            )
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Không thể trích xuất văn bản từ file '{extension or 'không rõ định dạng'}': {exc}"
        ) from exc

    text = text.strip()

    if not text:
        raise ValueError("Tài liệu không chứa văn bản trích xuất được.")

    if len(text) > MAX_DOCUMENT_CHARS:
        raise ValueError(
            f"Văn bản trích xuất vượt quá giới hạn {MAX_DOCUMENT_CHARS:,} ký tự. "
            "Vui lòng chia nhỏ tài liệu."
        )

    return text
